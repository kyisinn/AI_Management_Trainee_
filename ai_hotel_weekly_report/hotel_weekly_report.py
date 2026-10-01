#!/usr/bin/env python3
"""Create weekly hotel performance messages from CSV or XLSX input."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import json
import os
import re
import smtplib
import statistics
import sys
import urllib.error
import urllib.request
from email.message import EmailMessage
from pathlib import Path
from typing import Any


REQUIRED_COLUMNS = ("brandname", "revenue", "Order", "Boss Gmail")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def read_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            import openpyxl  # type: ignore
        except ImportError as exc:
            raise RuntimeError("Reading XLSX requires openpyxl: python3 -m pip install openpyxl") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value).strip() if value is not None else "" for value in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:] if any(value is not None for value in row)]
    else:
        raise ValueError("Input must be .csv, .xlsx, or .xlsm")

    missing = [column for column in REQUIRED_COLUMNS if column not in (rows[0].keys() if rows else [])]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    return rows


def as_number(value: Any, field: str, row_number: int) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Row {row_number}: {field} must be numeric") from exc
    if number < 0:
        raise ValueError(f"Row {row_number}: {field} cannot be negative")
    return number


def normalize(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    clean: list[dict[str, Any]] = []
    for index, row in enumerate(rows, start=2):
        hotel = str(row.get("brandname") or "").strip()
        email = str(row.get("Boss Gmail") or "").strip()
        if not hotel:
            raise ValueError(f"Row {index}: brandname is required")
        if not EMAIL_RE.match(email):
            raise ValueError(f"Row {index}: invalid owner email: {email!r}")
        clean.append(
            {
                "hotel": hotel,
                "revenue": as_number(row.get("revenue"), "revenue", index),
                "orders": as_number(row.get("Order"), "Order", index),
                "owner_email": email,
            }
        )
    if not clean:
        raise ValueError("No hotel rows found")
    return clean


def previous_week(run_date: dt.date) -> tuple[dt.date, dt.date]:
    this_monday = run_date - dt.timedelta(days=run_date.weekday())
    end = this_monday - dt.timedelta(days=1)
    start = end - dt.timedelta(days=6)
    return start, end


def classify(rows: list[dict[str, Any]], run_date: dt.date, test_email: str | None) -> list[dict[str, Any]]:
    median_revenue = statistics.median(row["revenue"] for row in rows)
    median_orders = statistics.median(row["orders"] for row in rows)
    week_start, week_end = previous_week(run_date)
    output = []
    for row in rows:
        doing_well = row["revenue"] >= median_revenue and row["orders"] >= median_orders
        no_activity = row["revenue"] == 0 and row["orders"] == 0
        result = dict(row)
        result.update(
            {
                "status": "Doing well" if doing_well else "Needs improvement",
                "no_activity": no_activity,
                "median_revenue": median_revenue,
                "median_orders": median_orders,
                "revenue_vs_median": row["revenue"] - median_revenue,
                "orders_vs_median": row["orders"] - median_orders,
                "week_start": week_start.isoformat(),
                "week_end": week_end.isoformat(),
                "recipient": test_email or row["owner_email"],
                "test_mode": bool(test_email),
            }
        )
        output.append(result)
    return output


def fallback_message(row: dict[str, Any]) -> dict[str, str]:
    hotel = html.escape(row["hotel"])
    revenue = f"{row['revenue']:,.0f}"
    orders = f"{row['orders']:,.0f}"
    median_revenue = f"{row['median_revenue']:,.0f}"
    median_orders = f"{row['median_orders']:,.0f}"
    if row["status"] == "Doing well":
        lead = "You finished the week at or above the hotel group median on both revenue and bookings."
        action = "Keep the campaigns and availability settings that supported this result, and review the strongest booking sources for repeatable wins."
    elif row["no_activity"]:
        lead = "No revenue or bookings were recorded for the week."
        action = "Please first confirm that the data feed is complete. If it is, review availability, pricing, and active sales channels before the next reporting cycle."
    else:
        weak = []
        if row["revenue"] < row["median_revenue"]:
            weak.append("revenue")
        if row["orders"] < row["median_orders"]:
            weak.append("bookings")
        lead = f"This week, {html.escape(' and '.join(weak))} finished below the hotel group median."
        action = "Review availability and pricing for the coming two weeks, then focus promotions on the channels that produced the most recent bookings."
    body = (
        f"<p>Hello {hotel} team,</p>"
        f"<p>Here is your performance summary for {row['week_start']} to {row['week_end']}. "
        f"You recorded <strong>{orders} bookings</strong> and <strong>{revenue}</strong> in revenue.</p>"
        f"<p>{lead} The comparison benchmark was {median_orders} bookings and {median_revenue} in revenue.</p>"
        f"<p><strong>Recommended next step:</strong> {action}</p>"
        "<p>Best regards,<br>Hotel Partnerships Team</p>"
    )
    return {"subject": f"Weekly performance summary - {row['hotel']}", "html_body": body}


def openai_message(row: dict[str, Any], api_key: str, model: str) -> dict[str, str]:
    schema = {
        "type": "object",
        "properties": {
            "subject": {"type": "string"},
            "html_body": {"type": "string"},
        },
        "required": ["subject", "html_body"],
        "additionalProperties": False,
    }
    instructions = (
        "Write a concise weekly hotel performance email. Use only the supplied facts. "
        "Do not invent causes, targets, percentages, or comparisons. Be supportive and specific, "
        "include one practical next step, and produce safe simple HTML with no links or images."
    )
    payload = {
        "model": model,
        "store": False,
        "instructions": instructions,
        "input": json.dumps(row, ensure_ascii=False),
        "text": {
            "format": {
                "type": "json_schema",
                "name": "hotel_weekly_email",
                "strict": True,
                "schema": schema,
            }
        },
    }
    request = urllib.request.Request(
        "https://api.openai.com/v1/responses",
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        data = json.loads(response.read().decode("utf-8"))
    output_text = data.get("output_text")
    if not output_text:
        for item in data.get("output", []):
            for content in item.get("content", []):
                if content.get("type") == "output_text":
                    output_text = content.get("text")
                    break
    message = json.loads(output_text or "{}")
    if not isinstance(message.get("subject"), str) or not isinstance(message.get("html_body"), str):
        raise ValueError("OpenAI response did not contain subject and html_body")
    return message


def send_email(row: dict[str, Any], message: dict[str, str]) -> None:
    required = ["SMTP_HOST", "SMTP_USERNAME", "SMTP_PASSWORD", "SMTP_FROM"]
    missing = [name for name in required if not os.getenv(name)]
    if missing:
        raise RuntimeError(f"Missing SMTP settings: {', '.join(missing)}")
    email = EmailMessage()
    email["From"] = os.environ["SMTP_FROM"]
    email["To"] = row["recipient"]
    email["Subject"] = message["subject"]
    email.set_content("This report is available as an HTML email.")
    email.add_alternative(message["html_body"], subtype="html")
    with smtplib.SMTP(os.environ["SMTP_HOST"], int(os.getenv("SMTP_PORT", "587")), timeout=30) as smtp:
        smtp.starttls()
        smtp.login(os.environ["SMTP_USERNAME"], os.environ["SMTP_PASSWORD"])
        smtp.send_message(email)


def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_")
    return cleaned or "hotel"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--run-date", type=dt.date.fromisoformat, default=dt.date.today())
    parser.add_argument("--test-email", help="Redirect every message to this address")
    parser.add_argument("--use-ai", action="store_true")
    parser.add_argument("--send", action="store_true")
    args = parser.parse_args()

    if args.test_email and not EMAIL_RE.match(args.test_email):
        parser.error("--test-email must be a valid email address")
    if args.send and not args.test_email:
        parser.error("For safety, --send requires --test-email")

    rows = classify(normalize(read_rows(args.input)), args.run_date, args.test_email)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    previews = args.output_dir / "email_previews"
    previews.mkdir(exist_ok=True)
    api_key = os.getenv("OPENAI_API_KEY")
    model = os.getenv("OPENAI_MODEL", "gpt-5-mini")
    messages = []

    for row in rows:
        source = "deterministic_fallback"
        error = None
        message = fallback_message(row)
        if args.use_ai:
            if not api_key:
                error = "OPENAI_API_KEY is not set"
            else:
                try:
                    message = openai_message(row, api_key, model)
                    source = "openai"
                except (ValueError, urllib.error.URLError, TimeoutError) as exc:
                    error = str(exc)
        record = {**row, **message, "message_source": source, "ai_error": error}
        messages.append(record)
        (previews / f"{safe_filename(row['hotel'])}.html").write_text(message["html_body"], encoding="utf-8")
        if args.send:
            send_email(row, message)

    with (args.output_dir / "messages.json").open("w", encoding="utf-8") as handle:
        json.dump(messages, handle, ensure_ascii=False, indent=2)
    result_fields = [
        "hotel", "revenue", "orders", "status", "median_revenue", "median_orders",
        "revenue_vs_median", "orders_vs_median", "week_start", "week_end",
        "owner_email", "recipient", "test_mode", "message_source", "ai_error",
    ]
    with (args.output_dir / "results.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=result_fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(messages)
    print(f"Processed {len(messages)} hotels. Output: {args.output_dir}")
    print("Email mode:", "SENT TO TEST ADDRESS" if args.send else "DRY RUN")
    return 0


if __name__ == "__main__":
    sys.exit(main())

